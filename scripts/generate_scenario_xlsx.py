#!/usr/bin/env python3
"""Generuje SCENARIUS_TESTOW_RECZNY_WYMAGANIA.xlsx z danych scenariusza (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SCENARIUS_TESTOW_RECZNY_WYMAGANIA.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WYNIK_OPTIONS = '"PASS,FAIL,CZĘŚCIOWO,N/D,BLOK"'


def main() -> None:
    wb = Workbook()

    # --- Arkusz: Instrukcja ---
    ws0 = wb.active
    ws0.title = "Instrukcja"
    intro = [
        ("Scenariusz testów WMS — wypełnianie ręczne", True),
        ("Źródło wymagań: PDF „Wymagania_WMS_pelne_z_brygadzista” + SCENARIUS_TESTOW_RECZNY_WYMAGANIA.md", False),
        ("", False),
        ("W kolumnie „Wynik” wpisuj: PASS, FAIL, CZĘŚCIOWO, N/D (nie dotyczy), BLOK (brak danych/dostępu).", False),
        ("Możesz wybrać wartość z listy rozwijanej w arkuszu „Scenariusz”.", False),
        ("", False),
        ("Mapowanie ról PDF → kod:", True),
        ("Administrator → ADMIN", False),
        ("Kierownik → MANAGER", False),
        ("Brygadzista → FOREMAN", False),
        ("Magazynier → WORKER", False),
        ("", False),
        ("Data testów:", False),
        ("Środowisko (localhost / docker / serwer):", False),
        ("Wersja / commit:", False),
        ("Tester:", False),
    ]
    for i, (text, bold) in enumerate(intro, start=1):
        c = ws0.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 85

    # --- Arkusz: Scenariusz (główna tabela) ---
    ws = wb.create_sheet("Scenariusz", 1)
    headers = [
        "Obszar",
        "ID wymagania",
        "Nazwa wymagania (skrót)",
        "Id kroku",
        "Krok / czynność",
        "Oczekiwany wynik",
        "Wynik",
        "Uwagi",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    rows: list[tuple[str, str, str, str, str, str]] = []
    # (obszar, req_id, req_title, step_id, krok, oczekiwane)

    def r(
        obs: str,
        rid: str,
        title: str,
        sid: str,
        krok: str,
        occ: str,
    ) -> None:
        rows.append((obs, rid, title, sid, krok, occ))

    # A — MUST funkcjonalne
    r("A Funkcjonalne MUST", "M-F1", "Uwierzytelnianie", "A1", "Wejdź na stronę logowania web", "Pola: login (5 cyfr) i hasło")
    r("A Funkcjonalne MUST", "M-F1", "Uwierzytelnianie", "A2", "Wpisz błędne hasło 2×", "Czytelny komunikat błędu, bez ujawniania „czy user istnieje”")
    r("A Funkcjonalne MUST", "M-F1", "Uwierzytelnianie", "A3", "Zaloguj poprawnie (np. 00001 + hasło admina)", "Przekierowanie na pulpit")
    r("A Funkcjonalne MUST", "M-F1", "Uwierzytelnianie", "A4", "(Mobilka) logowanie magazyniera", "Kod + hasło, wejście do aplikacji")

    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B1", "Jako ADMIN: Użytkownicy, Ustawienia", "Dostęp")
    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B2", "Jako WORKER: URL /users", "Brak dostępu / przekierowanie")
    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B3", "Jako WORKER: /settings", "Brak lub ograniczony dostęp")
    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B4", "Jako FOREMAN: Użytkownicy", "Brak pełnego dostępu jak ADMIN")
    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B5", "Jako MANAGER: słowniki vs użytkownicy", "Zgodnie z macierzą PDF (odczyt użytkowników)")
    r("A Funkcjonalne MUST", "M-F2", "RBAC", "B6", "API: token WORKER na endpoint ADMIN (Postman)", "HTTP 403 / odmowa")

    r("A Funkcjonalne MUST", "M-F3", "Kartoteka SKU", "C1", "Produkty: dodaj (SKU, nazwa, jednostka)", "Zapis, widoczny na liście")
    r("A Funkcjonalne MUST", "M-F3", "Kartoteka SKU", "C2", "Edytuj produkt", "Zapis, spójność danych")
    r("A Funkcjonalne MUST", "M-F3", "Kartoteka SKU", "C3", "Podgląd / wyszukiwanie po SKU", "Działa")

    r("A Funkcjonalne MUST", "M-F4", "Topologia magazynu", "D1", "Lokalizacje: dodaj hierarchię (rząd/regał/półka)", "Zapis bez błędu")
    r("A Funkcjonalne MUST", "M-F4", "Topologia magazynu", "D2", "Lista/drzewo lokalizacji", "Logiczna hierarchia")

    r("A Funkcjonalne MUST", "M-F5", "Prezentacja stanów", "E1", "Stany: produkt + lokalizacja", "Widać ilość zgodną z operacjami")
    r("A Funkcjonalne MUST", "M-F5", "Prezentacja stanów", "E2", "Po PZ/MM stan się zmienia", "Zgodność z oczekiwaniem")

    r("A Funkcjonalne MUST", "M-F6", "Rejestr ruchów", "F1", "Rejestr ruchów po PZ/MM", "Wpis: typ, czas, user, ilość")
    r("A Funkcjonalne MUST", "M-F6", "Rejestr ruchów", "F2", "Próba edycji/usunięcia wpisu z UI", "Brak możliwości (historia nieedytowalna)")

    r("A Funkcjonalne MUST", "M-F7", "PZ", "G1", "PZ: utwórz dokument i pozycje", "Zapis / workflow zgodny z UI")
    r("A Funkcjonalne MUST", "M-F7", "PZ", "G2", "Przyjęcie na strefę buforową", "Stany rosną")
    r("A Funkcjonalne MUST", "M-F7", "PZ", "G3", "Zakończenie procesu PZ", "Status końcowy, stany spójne")

    r("A Funkcjonalne MUST", "M-F8", "Put-away", "H1", "Rozmieszczanie: bufor → docelowa lokalizacja", "Stany zgodne (spadek w buforze, wzrost w docelowej)")
    r("A Funkcjonalne MUST", "M-F8", "Put-away", "H2", "Rejestr ruchów", "Wpis widoczny")

    r("A Funkcjonalne MUST", "M-F9", "MM", "I1", "MM: źródło i cel, walidacja", "Poprawne reguły biznesowe")
    r("A Funkcjonalne MUST", "M-F9", "MM", "I2", "Zatwierdź MM", "Stany przesunięte; ewent. zadania MOVE")

    r("A Funkcjonalne MUST", "M-F10", "RW / picking", "J1", "RW: kompletacja ze strefy picking", "Reguły (blokada bez stanu) działają")
    r("A Funkcjonalne MUST", "M-F10", "RW / picking", "J2", "Po kompletacji", "Stan maleje o ilość")
    r("A Funkcjonalne MUST", "M-F10", "RW / picking", "J3", "Kompletacja (web)", "Spójność z dokumentem")

    r("A Funkcjonalne MUST", "M-F11", "Inwentaryzacja", "K1", "Wprowadź stan faktyczny vs systemowy", "Widać różnice")
    r("A Funkcjonalne MUST", "M-F11", "Inwentaryzacja", "K2", "Zatwierdź korekty (role wg macierzy)", "Stany = faktyczne; Magazynier bez pełnych korekt wg PDF")

    r("A Funkcjonalne MUST", "M-F12", "Interfejs zarządczy", "L1", "Przeglądarka szeroka (laptop)", "Układ desktop (sidebar)")
    r("A Funkcjonalne MUST", "M-F12", "Interfejs zarządczy", "L2", "ADMIN: ustawienia, użytkownicy, słowniki", "Dostęp")

    r("A Funkcjonalne MUST", "M-F13", "UI mobilny (responsive)", "M1", "Wąskie okno / telefon", "Adaptacja, czytelne napisy, duże przyciski")
    r("A Funkcjonalne MUST", "M-F13", "UI mobilny (responsive)", "M2", "Zadania / skan", "Brak chaosu poziomym scrollowaniem")

    r("A Funkcjonalne MUST", "M-F14", "Lista zadań", "N1", "Zaloguj jako WORKER", "Widoczne Zadania")
    r("A Funkcjonalne MUST", "M-F14", "Lista zadań", "N2", "Filtry put-away, move, picking, inwent.", "Zgodne z dokumentami")
    r("A Funkcjonalne MUST", "M-F14", "Lista zadań", "N3", "Dokończ zadanie", "Status się zmienia")

    r("A Funkcjonalne MUST", "M-F15", "Tworzenie kont (admin)", "O1", "ADMIN: dodaj użytkownika (+ email)", "Konto + login 5 cyfr")
    r("A Funkcjonalne MUST", "M-F15", "Tworzenie kont (admin)", "O2", "SMTP wyłączone (dev)", "Komunikat / ścieżka testowa czy brak")
    r("A Funkcjonalne MUST", "M-F15", "Tworzenie kont (admin)", "O3", "Link ustawienia hasła", "Działa przy włączonym SMTP")

    # B — SHOULD funkcjonalne
    r("B Funkcjonalne SHOULD", "S-F1", "Generowanie zadań", "P1", "Po zatwierdzeniu PZ", "Zadanie put-away lub równoważny workflow")
    r("B Funkcjonalne SHOULD", "S-F3", "Statusy zapasu", "Q1", "Zmiana statusu (np. Dostępny/Zablokowany)", "Widać na liście / filtrach")
    r("B Funkcjonalne SHOULD", "S-F3", "Statusy zapasu", "Q2", "Ledger / audit", "Ślad operacji")
    r("B Funkcjonalne SHOULD", "S-F4", "Raportowanie", "R1", "Raporty: stany / ruchy", "Sensowne dane")
    r("B Funkcjonalne SHOULD", "S-F4", "Raportowanie", "R2", "Eksport CSV/PDF jeśli jest", "Pobranie pliku")
    r("B Funkcjonalne SHOULD", "S-F5", "Audit log", "S1", "Zmień atrybut produktu", "Wpis w dzienniku: kto, kiedy")
    r("B Funkcjonalne SHOULD", "S-F5", "Audit log", "S2", "FOREMAN vs WORKER dostęp", "Zgodnie z macierzą")
    r("B Funkcjonalne SHOULD", "WF1", "Poza zakresem", "-", "Brak integracji WooCommerce/Shopify", "N/D — nie testować integracji")

    # C — COULD
    r("C Funkcjonalne COULD", "C-F1", "PDF dokumentów", "-", "PZ/MM/RW → generuj PDF", "Plik do pobrania / druku")
    r("C Funkcjonalne COULD", "C-F2", "Packing", "-", "Weryfikacja po picking", "Ekran/proces zgodny z zakresem")
    r("C Funkcjonalne COULD", "C-F3", "Partie / FEFO", "-", "Partie, daty ważności", "Polamodeluje FIFO/FEFO")
    r("C Funkcjonalne COULD", "C-F4", "Zwroty", "-", "Proces zwrotu", "Rejestracja i ścieżki magazynowe")
    r("C Funkcjonalne COULD", "C-F5", "Multi-warehouse", "-", "Wiele magazynów w instancji", "Wyodrębnienie stanów")

    # D — Niefunkcjonalne MUST
    r("D Niefunkcjonalne MUST", "N-01", "Hasła hash", "T1", "Baza: pole hasła", "Wygląda na hash, nie plaintext")
    r("D Niefunkcjonalne MUST", "N-01", "Hasła hash", "T2", "Logowanie", "Tylko poprawne hasło")
    r("D Niefunkcjonalne MUST", "N-02", "HTTPS", "U1", "Produkcja/staging", "https, brak mixed content")
    r("D Niefunkcjonalne MUST", "N-02", "HTTPS", "U2", "Dev localhost", "http lokalnie OK")
    r("D Niefunkcjonalne MUST", "N-03", "Egzekwowanie ról API", "-", "Patrz B6", "403 przy obejściu UI")
    r("D Niefunkcjonalne MUST", "N-05", "Czas odpowiedzi", "V1", "Listy: zadania, dokumenty, stany (~10 prób)", "Subiektywnie większość < 2 s (95% wymaga load testu)")
    r("D Niefunkcjonalne MUST", "N-05", "Czas odpowiedzi", "V2", "Wyższe obciążenie", "Dokumentacja degradacji / benchmarki")
    r("D Niefunkcjonalne MUST", "N-06", "Atomowość dokumentów", "W1", "Błąd w trakcie zatwierdzania (kopia bazy!)", "Brak pośrednich stanów / rollback")
    r("D Niefunkcjonalne MUST", "N-08", "UI mobilny NFR", "-", "Patrz M-F13 + mobilka", "Duże cele, czytelność")
    r("D Niefunkcjonalne MUST", "N-11", "Rejestrowanie działań", "X1", "Login, PZ, zmiana stanu", "Audit z datą i użytkownikiem")
    r("D Niefunkcjonalne MUST", "N-13", "Jakość kodu", "Y1", "ruff / ESLint", "Akceptowalny raport zespołu")
    r("D Niefunkcjonalne MUST", "N-15", "Przeglądarki", "Z1", "Chrome", "Scenariusz OK")
    r("D Niefunkcjonalne MUST", "N-15", "Przeglądarki", "Z2", "Firefox", "Scenariusz OK")
    r("D Niefunkcjonalne MUST", "N-15", "Przeglądarki", "Z3", "Edge", "Scenariusz OK")

    # E — SHOULD NFR
    r("E Niefunkcjonalne SHOULD", "N-04", "Polityka haseł", "-", "Min. 8 znaków, litery+cyfry; rotacja 180 dni", "Walidacja / polityka w systemie")
    r("E Niefunkcjonalne SHOULD", "N-07", "Backup BDP", "-", "≥1/dzień, odtwarzanie", "Procedura ops / poza aplikacją")
    r("E Niefunkcjonalne SHOULD", "N-09", "Liczba kroków zadania", "-", "Od wyboru do końca ≤ ~5 ekranów", "Policz dla typowego zadania")
    r("E Niefunkcjonalne SHOULD", "N-10", "Spójność UI", "-", "Statusy, ikony, kolory spójne", "Przejście modułami")
    r("E Niefunkcjonalne SHOULD", "N-12", "Retencja logów", "-", "12 mies. + filtry data/user/typ", "Zapytania / ustawienia")
    r("E Niefunkcjonalne SHOULD", "N-14", "Dokumentacja", "-", "DEMO.md, README, Docker", "Opis architektury i startu")

    # F
    r("F Could/Won't NFR", "N-16", "Konteneryzacja", "-", "docker compose działa", "Uruchomienie środowiska")
    r("F Could/Won't NFR", "N-17", "Offline mobilka", "-", "Tryb samolot", "Komunikat o braku sieci; brak pełnego offline wg PDF")

    row_idx = 2
    for obs, rid, title, sid, krok, occ in rows:
        ws.cell(row=row_idx, column=1, value=obs)
        ws.cell(row=row_idx, column=2, value=rid)
        ws.cell(row=row_idx, column=3, value=title)
        ws.cell(row=row_idx, column=4, value=sid)
        ws.cell(row=row_idx, column=5, value=krok)
        ws.cell(row=row_idx, column=6, value=occ)
        ws.cell(row=row_idx, column=7, value=None)  # Wynik — do wypełnienia
        ws.cell(row=row_idx, column=8, value=None)  # Uwagi
        for c in range(1, 9):
            ws.cell(row=row_idx, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        row_idx += 1

    widths = [22, 14, 28, 10, 42, 38, 14, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_idx - 1}"

    dv = DataValidation(
        type="list",
        formula1=WYNIK_OPTIONS,
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Wybierz z listy lub zostaw puste."
    dv.errorTitle = "Wynik"
    last_data = row_idx - 1
    dv.add(f"G2:G{last_data}")
    ws.add_data_validation(dv)

    # --- Arkusz: Podsumowanie ---
    ws2 = wb.create_sheet("Podsumowanie")
    summary = [
        ("Liczba kroków w scenariuszu", str(len(rows))),
        ("", ""),
        ("Data zakończenia testów", ""),
        ("Środowisko", ""),
        ("Commit / wersja", ""),
        ("Liczba FAIL", ""),
        ("Liczba BLOK", ""),
        ("Uwagi ogólne (top 3 defekty)", ""),
    ]
    for i, (k, v) in enumerate(summary, start=1):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v if v else None)
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 55

    wb.save(OUT)
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    main()
