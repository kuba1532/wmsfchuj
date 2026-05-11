#!/usr/bin/env python3
"""
Generuje prezentację PPTX (krok po kroku: web + mobilka) dla promotora.

Zależność: pip install python-pptx

Uruchomienie z katalogu głównego repo:
  python3 scripts/generate_promoter_presentation_pptx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    print("Zainstaluj: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "prezentacja_demo" / "WMS_Promotor_Web_i_Mobilka.pptx"
ZRZUTY = ROOT / "prezentacja_demo" / "zrzuty"
E2E = ROOT / "e2e_tests" / "screens"

# (ścieżka względem ZRZUTY lub bezwzględna, tytuł slajdu, podtytuł/opis)
SLIDES: list[tuple[Path, str, str]] = [
    (
        ZRZUTY / "01_login.png",
        "1. Panel web — logowanie",
        "Kod 5 cyfr + hasło; dostęp do modułów wg roli (RBAC).",
    ),
    (
        ZRZUTY / "02_dashboard.png",
        "2. Pulpit",
        "Podsumowanie: produkty, zadania, dokumenty; szybki start procesów magazynowych.",
    ),
    (
        ZRZUTY / "03_pz_lista.png",
        "3. Przyjęcia (PZ) — lista",
        "Dokumenty przyjęć ze statusami; kolejne kroki: rejestracja na bufor, rozmieszczenie.",
    ),
    (
        ZRZUTY / "04_pz_form.png",
        "4. Nowe PZ — formularz",
        "Wybór dostawcy, pozycji, lokalizacji odłożenia; utworzenie dokumentu.",
    ),
    (
        ZRZUTY / "05_pz_lista_nowy.png",
        "5. PZ utworzone",
        "Nowy dokument na liście; przejście do rejestracji na strefie przyjęć.",
    ),
    (
        ZRZUTY / "06_pz_w_trakcie.png",
        "6. PZ w trakcie",
        "Po rejestracji: towar na buforze, zadania put-away gotowe do realizacji.",
    ),
    (
        ZRZUTY / "07_putaway_lista.png",
        "7. Rozmieszczanie (put-away) — lista",
        "Zadania przeniesienia z bufora na lokalizacje docelowe.",
    ),
    (
        ZRZUTY / "08_putaway_stepper.png",
        "8. Krok po kroku odłożenia",
        "Proces dla magazyniera: pobranie → przeniesienie → potwierdzenie.",
    ),
    (
        ZRZUTY / "09_pz_zakonczony.png",
        "9. PZ zakończone",
        "Dokument zamknięty po zrealizowaniu odłożeń; spójność ze stanami.",
    ),
    (
        ZRZUTY / "10_mm_lista.png",
        "10. Przesunięcia (MM)",
        "Przesunięcia wewnętrzne między lokalizacjami; zatwierdzenie i zadania MOVE.",
    ),
    (
        ZRZUTY / "11_rw_lista.png",
        "11. Wydania (RW)",
        "Wydanie z jawnej lokalizacji pobrania i odbiorcy ze słownika; zatwierdzenie i kompletacja.",
    ),
    (
        ZRZUTY / "12_stock.png",
        "12. Stany magazynowe",
        "Widok ilości per produkt i lokalizacja (np. AVAILABLE / BLOCKED).",
    ),
    (
        ZRZUTY / "13_ledger.png",
        "13. Rejestr ruchów",
        "Nieedytowalna historia ruchów powiązana z operacjami.",
    ),
    (
        ZRZUTY / "14_audit.png",
        "14. Dziennik zdarzeń (audit)",
        "Kto, kiedy, co zrobił w systemie — wsparcie dla nadzoru i zgodności.",
    ),
    (
        ZRZUTY / "15_raporty.png",
        "15. Raporty",
        "Eksporty (np. CSV/XLSX) dla analiz i zestawień.",
    ),
    (
        ZRZUTY / "16_users.png",
        "16. Użytkownicy i role",
        "Administracja kontami; macierz uprawnień (ADMIN, MANAGER, FOREMAN, WORKER).",
    ),
    (
        E2E / "flutter_worker_login.png",
        "17. Aplikacja mobilna (Flutter) — logowanie",
        "Ten sam backend API; na fizycznym telefonie adres serwera = IP komputera (nie 127.0.0.1).",
    ),
    (
        E2E / "flutter_worker_round2.png",
        "18. Aplikacja mobilna — praca w terenie",
        "Zadania, PZ, MM, RW, podgląd stanu; spójność z panelem webowym.",
    ),
]


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    sub = slide.placeholders[1]
    sub.text = subtitle
    for p in sub.text_frame.paragraphs:
        p.font.size = Pt(18)


def _add_image_slide(prs: Presentation, image: Path, title: str, caption: str) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    margin = Inches(0.45)
    title_h = Inches(0.85)
    cap_h = Inches(0.95)
    usable_w = prs.slide_width - 2 * margin

    tb = slide.shapes.add_textbox(margin, margin, usable_w, title_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)

    img_top = margin + title_h + Inches(0.08)
    img_bottom_reserve = cap_h + Inches(0.35)
    max_img_h = prs.slide_height - img_top - img_bottom_reserve
    max_img_w = usable_w

    if not image.is_file():
        cap = slide.shapes.add_textbox(margin, img_top, usable_w, Inches(2))
        cap.text_frame.paragraphs[0].text = f"[Brak pliku: {image.name}]"
        cap.text_frame.paragraphs[0].font.size = Pt(14)
    else:
        slide.shapes.add_picture(str(image), margin, img_top, width=max_img_w, height=max_img_h)

    cap_top = prs.slide_height - margin - cap_h
    cb = slide.shapes.add_textbox(margin, cap_top, usable_w, cap_h)
    ctf = cb.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = caption
    cp.font.size = Pt(14)
    cp.alignment = PP_ALIGN.LEFT


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(
        prs,
        "System WMS — demonstracja procesu",
        "Panel web + aplikacja mobilna (Flutter)\nKrok po kroku — materiał dla promotora\n"
        + str(ROOT.name),
    )

    _add_title_slide(
        prs,
        "Zakres prezentacji",
        "• Web: logowanie → PZ (przyjęcie, put-away) → MM → RW → stany → rejestr → audit → raporty → użytkownicy\n"
        "• Mobilka: logowanie i realizacja zadań / dokumentów przy tym samym API\n"
        "Szczegółowy scenariusz live: prezentacja_demo/SKRYPT_DEMO.md",
    )

    missing = []
    for path, title, cap in SLIDES:
        if not path.is_file():
            missing.append(path)
        _add_image_slide(prs, path, title, cap)

    if missing:
        print("Uwaga — brakujące pliki (slajdy zastąpione komunikatem):", file=sys.stderr)
        for m in missing:
            print(f"  {m}", file=sys.stderr)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    main()
