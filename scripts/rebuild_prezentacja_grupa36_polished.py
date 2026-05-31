#!/usr/bin/env python3
"""
Odbudowuje końcowe slajdy prezentacji grupa36 w konwencji istniejących slajdów
(duplikat slajdu „Testy” — ten sam nagłówek/stopka/numeracja, czysta treść + obrazy).

Uruchom po przywróceniu kopii 13-slajdowej (prezentacja grupa36_przed_uzupelnieniem.pptx).
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
BACKUP = Path("/Users/jakub/Downloads/prezentacja grupa36_przed_uzupelnieniem.pptx")
OUT_DOWNLOADS = Path("/Users/jakub/Downloads/prezentacja grupa36.pptx")
OUT_REPO = ROOT / "prezentacja_demo/PREZENTACJA_GRUPA36_uzupelniona.pptx"

# Slajd 8 (indeks 7) = „Testy…” — ma spójny nagłówek, lead, stopkę i numer strony.
TEMPLATE_SLIDE_INDEX = 7

WEB_PAIRS = [
    (
        "Funkcjonalności — widok w aplikacji webowej (1/2)",
        "Nawiązanie do slajdu „Funkcjonalności systemu”: dashboard oraz moduł dokumentów PZ w realnym panelu.",
        ("Dashboard", "Dokumenty PZ"),
        [
            ROOT / "prezentacja_demo/zrzuty/02_dashboard.png",
            ROOT / "prezentacja_demo/zrzuty/03_pz_lista.png",
        ],
    ),
    (
        "Funkcjonalności — widok w aplikacji webowej (2/2)",
        "Stany magazynowe i dokumenty MM — te same obszary z karty funkcji, pokazane w interfejsie użytkownika.",
        ("Stany magazynowe", "Dokumenty MM"),
        [
            ROOT / "prezentacja_demo/zrzuty/12_stock.png",
            ROOT / "prezentacja_demo/zrzuty/10_mm_lista.png",
        ],
    ),
]

USER_MOBILE_SHOT = Path(
    "/Users/jakub/.cursor/projects/Users-jakub-Downloads-WMS-2/assets/"
    "Zrzut_ekranu_2026-05-11_o_21.38.47-343536be-4ccc-4187-82f1-dd2d63ad5a6f.png"
)
MOBILE_SECOND = ROOT / "e2e_tests/screens/flutter_worker_round2.png"

ROADMAP_LINES = [
    "Wdrożenie produkcyjne: Docker, reverse proxy (HTTPS), secrets, backup i procedury rollback.",
    "CI/CD: lint + testy API + Vitest, artefakty buildów, tagowanie release.",
    "Real-time / sync: WebSocket lub SSE albo inteligentny polling per widok (bez „odświeżania na pale”).",
    "Integracje: rozszerzone CSV/XLSX, ewentualnie endpointy pod ERP (stany, potwierdzenia wysyłek).",
    "Multi-magazyn / multi-site: izolacja danych, filtry globalne, raporty skonsolidowane.",
    "TMS / logistyka (opcjonalnie): etykiety, śledzenie przesyłek — po decyzji biznesowej.",
    "Bezpieczeństwo: MFA dla ról admin, polityki haseł, audyt konfiguracji, rotacja kluczy JWT.",
    "Monitoring: metryki (np. Prometheus/Grafana), alerty na 5xx i kolejki zadań.",
    "Wydajność: indeksy, cache raportów, kolejki dla operacji masowych (skalowanie >~25 równoległych użytkowników).",
]

TEAM_LINES = [
    "Jakub Rzepkowski — backend (FastAPI, API, migracje, dokumenty PZ/MM/RW, zadania, RBAC) oraz aplikacja mobilna Flutter (worker, API, skaner, iOS/Android).",
    "Miłosz Marek — frontend web (React, TypeScript, Vite, MUI): moduły operacyjne, formularze dokumentów, stany, zadania, UX i integracja z API.",
    "Grzegorz Pawlak — dokumentacja i jakość formalna: scenariusze testów manualnych, arkusze ryzyk/monitoringu, raporty UAT/GO–NO–GO, feedback interesariuszy oraz materiały porównawcze dla promotora/klienta.",
]


def _duplicate_slide_shapes(prs: Presentation, source_index: int):
    """Nowy slajd: puste layout[0] + deepcopy wszystkich kształtów ze wzorca."""
    source = prs.slides[source_index]
    dest = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in source.shapes:
        dest.shapes._spTree.insert_element_before(copy.deepcopy(shape.element), "p:extLst")
    return dest


def _trim_tests_slide_to_header_footer(slide) -> None:
    """Zostaw tylko nagłówek (0), lead (1), stopkę i numer strony (ostatnie dwa kształty)."""
    n = len(slide.shapes)
    # Na slajdzie „Testy” środek to indeksy 2..18 (patrz audyt bbox).
    for i in range(n - 1, 1, -1):
        if i <= 18:
            el = slide.shapes[i].element
            el.getparent().remove(el)


def _set_title_lead(slide, title: str, lead: str) -> None:
    if slide.shapes[0].has_text_frame:
        slide.shapes[0].text_frame.text = title
    if slide.shapes[1].has_text_frame:
        slide.shapes[1].text_frame.text = lead


def _style_photo_card(shape) -> None:
    """Delikatna „karta” pod zrzut — spójna z tłem slajdów merytorycznych."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    shape.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    shape.line.width = Pt(1)


def _add_accent_strip(slide, left, top, height) -> None:
    """Cienki pasek w kolorze akcentu (jak zielenie na slajdzie Testy)."""
    w = Inches(0.07)
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x10, 0xB9, 0x81)
    sh.line.fill.background()


def _add_figure_caption(slide, left, top, width, text: str) -> None:
    box = slide.shapes.add_textbox(left, top, width, Inches(0.28))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.name = "Aptos"
    p.font.italic = True
    p.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def _set_page_number(slide, page: int) -> None:
    """Numer strony: mały kształt tekstowy z samą cyfrą (jak na wzorcu „Testy”)."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and len(t) <= 2:
            sh.text_frame.text = str(page)
            return
    raise RuntimeError("Nie znaleziono kształtu numeru strony na slajdzie wzorcowym.")


def _add_two_web_screenshots(
    slide,
    left_path: Path,
    right_path: Path,
    captions: tuple[str, str],
) -> None:
    """Dwa zrzuty w „kartach” + podpis — najpierw tło (na spód), potem obrazy."""
    y_card = Inches(1.02)
    h_card = Inches(3.52)
    gap = Inches(0.22)
    card_w = Inches(4.48)
    x_left = Inches(0.48)
    x_right = x_left + card_w + gap
    inset = Inches(0.12)
    pic_w = card_w - 2 * inset

    for x in (x_left, x_right):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_card, card_w, h_card)
        _style_photo_card(card)
        _add_accent_strip(slide, x + Inches(0.06), y_card + Inches(0.12), h_card - Inches(0.24))

    y_pic = y_card + inset
    slide.shapes.add_picture(str(left_path), x_left + inset, y_pic, width=pic_w)
    slide.shapes.add_picture(str(right_path), x_right + inset, y_pic, width=pic_w)

    cap_y = y_card + h_card + Inches(0.06)
    _add_figure_caption(slide, x_left, cap_y, card_w, captions[0])
    _add_figure_caption(slide, x_right, cap_y, card_w, captions[1])


def _add_two_mobile_screenshots_stacked(slide, top_path: Path, bottom_path: Path) -> None:
    """Mobilka: dwie karty jedna pod drugą + stała wysokość obrazu w środku karty."""
    x_card = Inches(0.95)
    card_w = Inches(8.2)
    h_card = Inches(2.78)
    y1 = Inches(1.02)
    gap = Inches(0.14)
    y2 = y1 + h_card + gap
    inset = Inches(0.14)
    pic_h = h_card - 2 * inset

    for y in (y1, y2):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_card, y, card_w, h_card)
        _style_photo_card(card)
        _add_accent_strip(slide, x_card + Inches(0.06), y + Inches(0.12), h_card - Inches(0.24))

    slide.shapes.add_picture(str(top_path), x_card + inset, y1 + inset, height=pic_h)
    slide.shapes.add_picture(str(bottom_path), x_card + inset, y2 + inset, height=pic_h)

    cap_y1 = y1 + h_card + Inches(0.04)
    cap_y2 = y2 + h_card + Inches(0.04)
    _add_figure_caption(slide, x_card, cap_y1, card_w, "Logowanie / konfiguracja API na urządzeniu")
    _add_figure_caption(slide, x_card, cap_y2, card_w, "Lista zadań magazyniera (worker)")


def _add_bullet_block(slide, lines: list[str]) -> None:
    """Blok treści w miejscu „zielonego” obszaru ze slajdu Testy (w przybliżeniu bbox shape 4)."""
    left, top, width, height = Inches(0.55), Inches(1.05), Inches(9.0), Inches(5.55)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    if not lines:
        return
    tf.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        pr = tf.add_paragraph()
        pr.text = line
        pr.level = 0
    for pr in tf.paragraphs:
        pr.font.size = Pt(11)
        pr.font.name = "Aptos"
        pr.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        pr.space_after = Pt(5)


def main() -> None:
    if not BACKUP.is_file():
        raise SystemExit(f"Brak kopii 13-slajdowej: {BACKUP}")

    missing = []
    for _, _, _, paths in WEB_PAIRS:
        for img_path in paths:
            if not img_path.is_file():
                missing.append(img_path)
    if not USER_MOBILE_SHOT.is_file():
        missing.append(USER_MOBILE_SHOT)
    if not MOBILE_SECOND.is_file():
        missing.append(MOBILE_SECOND)
    if missing:
        raise SystemExit("Brak plików PNG:\n" + "\n".join(str(p) for p in missing))

    shutil.copy2(BACKUP, BACKUP.with_name(BACKUP.name + ".bak_polish"))
    prs = Presentation(str(BACKUP))
    if len(prs.slides) != 13:
        raise SystemExit(f"Oczekiwano 13 slajdów w kopii, jest: {len(prs.slides)}")

    start_page = 14

    for i, (title, lead, captions, paths) in enumerate(WEB_PAIRS):
        sl = _duplicate_slide_shapes(prs, TEMPLATE_SLIDE_INDEX)
        _trim_tests_slide_to_header_footer(sl)
        _set_title_lead(sl, title, lead)
        _set_page_number(sl, start_page + i)
        _add_two_web_screenshots(sl, paths[0], paths[1], captions)

    m_sl = _duplicate_slide_shapes(prs, TEMPLATE_SLIDE_INDEX)
    _trim_tests_slide_to_header_footer(m_sl)
    _set_title_lead(
        m_sl,
        "Proces magazynowy — widok w aplikacji mobilnej (Flutter)",
        "Nawiązanie do slajdu „Proces działania WMS”: ta sama logika operacji, widziana z perspektywy magazyniera na telefonie.",
    )
    _set_page_number(m_sl, start_page + len(WEB_PAIRS))
    _add_two_mobile_screenshots_stacked(m_sl, USER_MOBILE_SHOT, MOBILE_SECOND)

    r_sl = _duplicate_slide_shapes(prs, TEMPLATE_SLIDE_INDEX)
    _trim_tests_slide_to_header_footer(r_sl)
    _set_title_lead(
        r_sl,
        "Plan rozwoju — od pilotażu do produkcji (uzupełnienie)",
        "Rozszerzenie slajdu „Plan rozwoju aplikacji”: konkrety techniczne i organizacyjne po pierwszym wdrożeniu.",
    )
    _set_page_number(r_sl, start_page + len(WEB_PAIRS) + 1)
    _add_bullet_block(r_sl, ROADMAP_LINES)

    t_sl = _duplicate_slide_shapes(prs, TEMPLATE_SLIDE_INDEX)
    _trim_tests_slide_to_header_footer(t_sl)
    _set_title_lead(
        t_sl,
        "Podział prac w zespole projektowym",
        "Uzupełnienie slajdu „Metodologia pracy zespołu”: kto odpowiadał za które warstwy systemu i dokumentację.",
    )
    _set_page_number(t_sl, start_page + len(WEB_PAIRS) + 2)
    _add_bullet_block(t_sl, TEAM_LINES)

    prs.save(str(OUT_DOWNLOADS))
    shutil.copy2(OUT_DOWNLOADS, OUT_REPO)
    print(f"Zapisano: {OUT_DOWNLOADS}")
    print(f"Kopia w repo: {OUT_REPO}")
    print(f"Łącznie slajdów: {len(prs.slides)}")


if __name__ == "__main__":
    main()
