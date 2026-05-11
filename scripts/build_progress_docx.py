#!/usr/bin/env python3
"""
Prosta konwersja Markdown -> DOCX (wystarczająca do dokumentu oddaniowego).
Nie renderuje tabel jako natywnych tabel DOCX, ale zachowuje pełną treść.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "SPRAWOZDANIE_POSTEPU_DPZ_WMS_2026-05-11.md"
OUT = ROOT / "SPRAWOZDANIE_POSTEPU_DPZ_WMS_2026-05-11.docx"


def main() -> None:
    md = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw in md:
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith(("1. ", "2. ", "3. ", "4. ", "5. ", "6. ", "7. ", "8. ", "9. ")):
            # Prosta obsługa list numerowanych
            doc.add_paragraph(line, style="List Number")
        else:
            doc.add_paragraph(line)

    doc.save(str(OUT))
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    main()
