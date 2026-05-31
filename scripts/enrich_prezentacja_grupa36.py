#!/usr/bin/env python3
"""
DEPRECATED: pierwsza wersja „wklejania” slajdów (puste layouty, inna konwencja wizualna).

Użyj zamiast tego: `scripts/rebuild_prezentacja_grupa36_polished.py`
— duplikuje slajd „Testy” (nagłówek, stopka, numeracja) i wstawia zrzuty + treść spójnie z resztą decka.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = Path("/Users/jakub/Downloads/prezentacja grupa36.pptx")
BACKUP = Path("/Users/jakub/Downloads/prezentacja grupa36_przed_uzupelnieniem.pptx")
OUT = SRC

WEB_SHOTS = [
    ROOT / "prezentacja_demo/zrzuty/02_dashboard.png",
    ROOT / "prezentacja_demo/zrzuty/03_pz_lista.png",
    ROOT / "prezentacja_demo/zrzuty/12_stock.png",
    ROOT / "prezentacja_demo/zrzuty/10_mm_lista.png",
]
MOBILE_SHOTS = [
    ROOT / "e2e_tests/screens/flutter_worker_login.png",
    ROOT / "e2e_tests/screens/flutter_worker_round2.png",
]


def _add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(9.3), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(22)
    p.font.bold = True


def _add_caption(slide, text: str, top: float) -> None:
    box = slide.shapes.add_textbox(Inches(0.35), Inches(top), Inches(9.3), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.italic = True


def _add_two_images_row(slide, left_path: Path, right_path: Path, y_top: float) -> None:
    w = Inches(4.55)
    h = Inches(2.85)
    slide.shapes.add_picture(str(left_path), Inches(0.35), Inches(y_top), width=w, height=h)
    slide.shapes.add_picture(str(right_path), Inches(5.05), Inches(y_top), width=w, height=h)


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _add_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.85), Inches(9.1), Inches(6.2))
    tf = box.text_frame
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(14)
        p.space_after = Pt(6)


def main() -> None:
    missing = [p for p in WEB_SHOTS + MOBILE_SHOTS if not p.is_file()]
    if missing:
        raise SystemExit("Brak plików PNG:\n" + "\n".join(str(p) for p in missing))

    if not SRC.is_file():
        raise SystemExit(f"Brak pliku źródłowego: {SRC}")

    shutil.copy2(SRC, BACKUP)

    prs = Presentation(str(SRC))

    # --- Slajd: web cz. 1 ---
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    _add_title(s1, "Zrzuty ekranowe — aplikacja webowa (1/2)")
    _add_two_images_row(s1, WEB_SHOTS[0], WEB_SHOTS[1], 0.82)
    _add_caption(s1, "Dashboard oraz lista dokumentów PZ.", 3.75)

    # --- Slajd: web cz. 2 ---
    s2 = prs.slides.add_slide(prs.slide_layouts[0])
    _add_title(s2, "Zrzuty ekranowe — aplikacja webowa (2/2)")
    _add_two_images_row(s2, WEB_SHOTS[2], WEB_SHOTS[3], 0.82)
    _add_caption(s2, "Stany magazynowe oraz dokumenty MM.", 3.75)

    # --- Slajd: mobilka ---
    s3 = prs.slides.add_slide(prs.slide_layouts[0])
    _add_title(s3, "Zrzuty ekranowe — aplikacja mobilna (Flutter)")
    _add_two_images_row(s3, MOBILE_SHOTS[0], MOBILE_SHOTS[1], 0.82)
    _add_caption(s3, "Logowanie magazyniera oraz widok operacji / zadań (worker).", 3.75)

    roadmap = [
        "Wdrożenie produkcyjne: konteneryzacja (Docker), reverse proxy (HTTPS), zmienne środowiskowe i backup bazy.",
        "CI/CD i jakość: pipeline (lint + testy API + Vitest), artefakty buildów, wersjonowanie release.",
        "Powiadomienia w czasie rzeczywistym: WebSocket/SSE lub lekki polling per moduł — szybsza synchronizacja web↔mobile bez obciążania list.",
        "Integracje: eksport/import CSV/XLSX rozszerzony, ewentualnie API dla ERP (np. stany zamówień, potwierdzenia wysyłek).",
        "Magazyn wielooddziałowy / multi-site: oddzielne magazyny, filtry globalne, raporty skonsolidowane.",
        "Zaawansowany TMS / śledzenie przesyłek, etykiety transportowe (opcjonalnie po decyzji biznesowej).",
        "Bezpieczeństwo: MFA dla ról administracyjnych, polityki haseł, audyt konfiguracji, rotacja kluczy JWT.",
        "Monitoring operacyjny: metryki (Prometheus/Grafana lub SaaS), alerty na błędy 5xx i długość kolejek zadań.",
        "Skalowanie wydajnościowe: tuning zapytań, cache odczytów raportowych, kolejki dla operacji masowych (powyżej ~25 równoległych użytkowników).",
    ]
    _add_bullet_slide(prs, "Plan rozwoju aplikacji — kierunki na przyszłość (uzupełnienie)", roadmap)

    team = [
        "Jakub Rzepkowski — backend (FastAPI, API, migracje, logika dokumentów PZ/MM/RW, zadania, RBAC) oraz aplikacja mobilna Flutter (worker, integracja z API, skaner, iOS/Android).",
        "Miłosz Marek — frontend web (React, TypeScript, Vite, MUI): layout, moduły operacyjne, formularze dokumentów, stany, zadania, spójność UX i integracja z API.",
        (
            "Grzegorz Pawlak — dokumentacja projektowa i formalizacja jakości: scenariusze testów manualnych, "
            "arkusze ryzyk i monitorowania, raporty UAT/GO–NO–GO, zbieranie uwag od interesariuszy oraz "
            "sekcja porównawcza „WMS vs alternatywy” w materiałach dla promotora i klienta (bez dublowania prac programistycznych)."
        ),
    ]
    _add_bullet_slide(prs, "Podział prac w zespole projektowym", team)

    prs.save(str(OUT))
    print(f"Kopia zapasowa: {BACKUP}")
    print(f"Zapisano: {OUT}")
    print(f"Dodano slajdy: {len(prs.slides) - 13} (łącznie slajdów: {len(prs.slides)})")


if __name__ == "__main__":
    main()
