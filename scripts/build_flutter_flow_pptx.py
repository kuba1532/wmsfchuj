"""Profesjonalna prezentacja PPTX — moduł mobilny WMS (Flutter).

Przeznaczenie: obrona / prezentacja przed promotorem — czytelne zrzuty, narracja,
slajdy kontekstu (cel, architektura, agenda, podsumowanie).

Wymaga zrzutów z testu integracyjnego:
  cd wms_worker && flutter test integration_test/wms_flow_golden_test.dart -d macos

Wyjście: prezentacja_demo/WMS_Flutter_przeplyw_mobilka.pptx
"""
from __future__ import annotations

import struct
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

REPO = Path(__file__).resolve().parent.parent
SRC = REPO / "wms_worker" / "build" / "integration_screenshots"
OUT = REPO / "prezentacja_demo" / "WMS_Flutter_przeplyw_mobilka.pptx"

# Styl „defense-ready”: jasne tło, kontrastowy nagłówek (czytelne przy projektorze)
BG = RGBColor(0xFA, 0xFA, 0xFA)
PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
ACCENT = RGBColor(0x39, 0x49, 0xAB)
BODY = RGBColor(0x37, 0x47, 0x4F)
MUTED = RGBColor(0x78, 0x90, 0x9C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def read_png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as f:
        if f.read(8) != b"\x89PNG\r\n\x1a\n":
            return 856, 1852
        _len = struct.unpack(">I", f.read(4))[0]
        ctype = f.read(4)
        if ctype != b"IHDR":
            return 856, 1852
        data = f.read(13)
        w, h = struct.unpack(">II", data[:8])
        return w, h


def set_paragraph_text(p, text: str, *, size: int, bold=False, color: RGBColor = BODY, align=PP_ALIGN.LEFT):
    p.text = text
    p.alignment = align
    if p.runs:
        r0 = p.runs[0]
        r0.font.size = Pt(size)
        r0.font.bold = bold
        r0.font.color.rgb = color


def add_cover_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, int(sw * 0.012), sh)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT

    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, int(sh * 0.42))
    hero.line.fill.background()
    hero.fill.solid()
    hero.fill.fore_color.rgb = PRIMARY

    tb = s.shapes.add_textbox(Inches(0.65), Inches(0.55), Inches(11.8), Inches(1.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    set_paragraph_text(
        p,
        "System WMS — moduł mobilny dla operatora magazynu",
        size=40,
        bold=True,
        color=WHITE,
    )

    st = s.shapes.add_textbox(Inches(0.65), Inches(1.85), Inches(11.5), Inches(1.0))
    tf = st.text_frame
    p = tf.paragraphs[0]
    set_paragraph_text(
        p,
        "Aplikacja Flutter (iOS / Android / desktop): uwierzytelnianie, przyjęcia PZ ze skanem "
        "kodu, realizacja zadań magazynowych z potwierdzeniem lokalizacji.",
        size=20,
        color=RGBColor(0xC5, 0xCA, 0xE9),
    )

    body = s.shapes.add_textbox(Inches(0.65), Inches(3.15), Inches(11.8), Inches(3.2))
    tf = body.text_frame
    tf.word_wrap = True
    bullets = [
        "Materiał demonstracyjny: pełny przepływ operacyjny zarejestrowany automatycznie (test integracyjny).",
        "Backend: REST API (FastAPI), JWT, role (WORKER), spójny model dokumentów i zadań z panelem webowym.",
        "Na slajdach: realne ekrany — od logowania, przez utworzenie i rejestrację PZ, do zakończenia zadania odłożenia.",
    ]
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph_text(p, "•  " + b, size=17, color=BODY)
        p.space_after = Pt(10)

    foot = s.shapes.add_textbox(Inches(0.65), Inches(6.85), Inches(11.8), Inches(0.45))
    tf = foot.text_frame
    p = tf.paragraphs[0]
    set_paragraph_text(
        p,
        "Uzupełnienie pracy dyplomowej  •  [Imię Nazwisko]  •  [Rok / uczelnia]",
        size=13,
        color=MUTED,
        align=PP_ALIGN.LEFT,
    )
    s.notes_slide.notes_text_frame.text = (
        "Slajd otwierający: przedstawić w jednym zdaniu, że moduł mobilny realizuje ten sam model danych co "
        "system webowy i jest przeznaczony dla pracownika magazynu w terenie."
    )


def add_text_only_slide(prs, title: str, bullets: list[str], speaker: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG

    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.55), Inches(0.12), Inches(1.05))
    accent.line.fill.background()
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT

    tb = s.shapes.add_textbox(Inches(0.75), Inches(0.45), Inches(11.5), Inches(0.95))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    set_paragraph_text(p, title, size=32, bold=True, color=PRIMARY)

    bb = s.shapes.add_textbox(Inches(0.75), Inches(1.55), Inches(11.5), Inches(5.5))
    tf = bb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph_text(p, "•  " + b, size=20, color=BODY)
        p.space_after = Pt(12)
        p.level = 0

    if speaker:
        s.notes_slide.notes_text_frame.text = speaker


def add_screenshot_slide(
    prs,
    *,
    step: str,
    title: str,
    takeaway: str,
    bullets: list[str],
    image: str,
    speaker: str = "",
):
    path = SRC / image
    if not path.exists():
        return
    w_px, h_px = read_png_size(path)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG

    # Nagłówek: krok + tytuł
    head = s.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.42))
    tf = head.text_frame
    p = tf.paragraphs[0]
    set_paragraph_text(p, step, size=13, bold=True, color=ACCENT)
    p2 = tf.add_paragraph()
    set_paragraph_text(p2, title, size=26, bold=True, color=PRIMARY)

    take = s.shapes.add_textbox(Inches(0.5), Inches(0.78), Inches(12.3), Inches(0.38))
    tf = take.text_frame
    p = tf.paragraphs[0]
    set_paragraph_text(p, takeaway, size=15, color=MUTED)

    # Obraz: maks. szerokość i wysokość — wyśrodkowany (czytelność na projektorze)
    max_w = Inches(11.8)
    max_h = Inches(5.35)
    aspect = w_px / h_px
    img_w = max_w
    img_h = img_w / aspect
    if img_h > max_h:
        img_h = max_h
        img_w = img_h * aspect
    left = (sw - img_w) / 2
    top = Inches(1.22)
    s.shapes.add_picture(str(path), left, top, width=img_w)

    # Dolne punkty (2–3 zdania)
    foot_top = Inches(6.58)
    fb = s.shapes.add_textbox(Inches(0.55), foot_top, Inches(12.2), Inches(0.75))
    tf = fb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets[:3]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph_text(p, "→  " + b, size=14, color=BODY)
        p.space_after = Pt(4)

    if speaker:
        s.notes_slide.notes_text_frame.text = speaker


def main():
    steps = [
        {
            "step": "Ekran 1 z 15",
            "title": "Uwierzytelnianie operatora",
            "takeaway": "Ten sam mechanizm co w panelu WWW: kod 5-cyfrowy + hasło, sesja JWT.",
            "image": "wms_ops_01_logowanie.png",
            "bullets": [
                "Identyfikacja pracownika magazynu (rola WORKER).",
                "Adres API widoczny na ekranie — ułatwia konfigurację w środowisku deweloperskim i na urządzeniu.",
            ],
            "speaker": "Podkreślić spójność z backendem: jedno źródło prawdy dla użytkowników i uprawnień.",
        },
        {
            "step": "Ekran 2 z 15",
            "title": "Lista zadań magazynowych",
            "takeaway": "Główny ekran pracy: typ zadania, status, filtry, bezpieczne zakończenie ze skanem lokalizacji.",
            "image": "wms_ops_02_zadania_przed.png",
            "bullets": [
                "Widoczne zadania przypisane do zalogowanego operatora.",
                "Nazewnictwo przystępne dla użytkownika (np. „Odłożenie po przyjęciu” zamiast wyłącznie PUTAWAY).",
            ],
            "speaker": "Można wspomnieć o rozróżnieniu: „Rozpocznij” rezerwuje pracę, „Potwierdź” wiąże się z walidacją miejsca.",
        },
        {
            "step": "Ekran 3 z 15",
            "title": "Tworzenie dokumentu przyjęcia (PZ)",
            "takeaway": "Formularz mobilny odzwierciedla reguły biznesowe: dostawca, pozycje, miejsce docelowe odłożenia.",
            "image": "wms_ops_03_pz_formularz.png",
            "bullets": [
                "Przyjęcie przygotowane pod pracę „z łapy”: krótki opis procesu na ekranie.",
                "Kolejne kroki: wybór dostawcy → pozycje z katalogu (skan/wyszukiwanie) → utworzenie dokumentu.",
            ],
        },
        {
            "step": "Ekran 4 z 15",
            "title": "Powiązanie z dostawcą",
            "takeaway": "Katalog towarów jest filtrowany — tylko asortyment przypisany do wybranego dostawcy.",
            "image": "wms_ops_04_pz_po_wyborze_dostawcy.png",
            "bullets": [
                "Ogranicza błędy wyboru i jest zgodne z walidacją po stronie API.",
            ],
        },
        {
            "step": "Ekran 5 z 15",
            "title": "Wybór towaru — katalog i skan",
            "takeaway": "Operator może wyszukać produkt lub użyć ścieżki „jak po skanie” (wpis kodu / skaner).",
            "image": "wms_ops_05_pz_wybor_towaru_sheet.png",
            "bullets": [
                "Arkusz dolny (bottom sheet) — wzorzec znany z aplikacji mobilnych.",
            ],
        },
        {
            "step": "Ekran 6 z 15",
            "title": "Wprowadzenie kodu (SKU / EAN)",
            "takeaway": "Ten sam dialog obsługuje wpis ręczny i integrację ze skanerem — na zrzucie symulacja kodu PRD-1001.",
            "image": "wms_ops_06_pz_dialog_kodu_jak_skaner.png",
            "bullets": [
                "Po zatwierdzeniu następuje zapytanie do API i lista produktów spełniających kryteria.",
            ],
        },
        {
            "step": "Ekran 7 z 15",
            "title": "Potwierdzenie pozycji z katalogu",
            "takeaway": "Wybór wiersza = dodanie pozycji do dokumentu PZ.",
            "image": "wms_ops_07_pz_lista_po_skanie.png",
            "bullets": [
                "Widać nazwę, SKU i ewentualny EAN — transparentność dla operatora.",
            ],
        },
        {
            "step": "Ekran 8 z 15",
            "title": "Parametry pozycji: ilość i lokalizacja docelowa",
            "takeaway": "Wskazanie regału (np. STO-01) już na etapie przyjęcia — przygotowanie pod zadanie odłożenia.",
            "image": "wms_ops_08_pz_pozycja_gotowa.png",
            "bullets": [
                "Backend wymaga lokalizacji magazynowej (nie bufora przyjęć) jako celu odłożenia.",
            ],
        },
        {
            "step": "Ekran 9 z 15",
            "title": "Dokument w statusie „Nowy”",
            "takeaway": "Lista dokumentów pokazuje numer PZ i powiązane zadania (po ich utworzeniu).",
            "image": "wms_ops_09_pz_lista_z_draftem.png",
            "bullets": [
                "Przycisk „Zarejestruj” uruchamia księgowanie na bufor i generowanie zadań PUTAWAY.",
            ],
        },
        {
            "step": "Ekran 10 z 15",
            "title": "Po rejestracji przyjęcia",
            "takeaway": "Towar na buforze; utworzone zadania odłożenia — identyfikatory widoczne przy dokumencie.",
            "image": "wms_ops_10_pz_po_zarejestruj.png",
            "bullets": [
                "Spójność dokument ↔ zadania: widać, że proces przechodzi do realizacji fizycznej.",
            ],
            "speaker": "Tu warto powiedzieć: rejestracja jest atomowa po stronie serwera — jeden krok, spójny stan.",
        },
        {
            "step": "Ekran 11 z 15",
            "title": "Realizacja zadania — przed startem",
            "takeaway": "Filtrowanie po typie zadania; operator rozpoczyna pracę przyciskiem „Biorę zadanie”.",
            "image": "wms_ops_11_zadanie_putaway.png",
            "bullets": [
                "Lista sortowana chronologicznie — najnowsze zadania u góry.",
            ],
        },
        {
            "step": "Ekran 12 z 15",
            "title": "Zadanie w realizacji",
            "takeaway": "Status „w trakcie”; zakończenie wymaga potwierdzenia kodem miejsca docelowego.",
            "image": "wms_ops_12_zadanie_w_trakcie.png",
            "bullets": [
                "Ochrona przed przypadkowym zamknięciem bez wykonania czynności.",
            ],
        },
        {
            "step": "Ekran 13 z 15",
            "title": "Walidacja lokalizacji",
            "takeaway": "Kod musi zgadzać się z lokalizacją docelową zadania — analogia do skanu etykiety regału.",
            "image": "wms_ops_13_dialog_potwierdz_lokalizacje.png",
            "bullets": [
                "Zmniejsza ryzyko błędnego księgowania w systemie.",
            ],
        },
        {
            "step": "Ekran 14 z 15",
            "title": "Po zakończeniu — widok aktywnych",
            "takeaway": "Zadanie zakończone nie zaśmieca domyślnej listy (ukryte zakończone).",
            "image": "wms_ops_14_po_zakonczeniu_zadania.png",
            "bullets": [
                "Operator widzi tylko to, co wymaga dalszej pracy.",
            ],
        },
        {
            "step": "Ekran 15 z 15",
            "title": "Weryfikacja — ujawnienie zadań zakończonych",
            "takeaway": "Opcja pokazania zadań zamkniętych — potwierdzenie sukcesu i audytowalność z perspektywy użytkownika.",
            "image": "wms_ops_15_zadania_zakonczone_widoczne.png",
            "bullets": [
                "Status „Zakończone” widoczny po włączeniu przełącznika.",
            ],
        },
    ]

    missing = [s["image"] for s in steps if not (SRC / s["image"]).exists()]
    if missing:
        print("Brak plików PNG — uruchom:")
        print("  cd wms_worker && flutter test integration_test/wms_flow_golden_test.dart -d macos")
        print("Brakuje:", ", ".join(missing))
        return 1

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover_slide(prs)

    add_text_only_slide(
        prs,
        "Agenda",
        [
            "Kontekst i cel modułu mobilnego.",
            "Architektura integracji z systemem WMS.",
            "Demonstracja przepływu: PZ → rejestracja → zadanie odłożenia → potwierdzenie lokalizacji.",
            "Podsumowanie możliwości i obszarów rozwoju.",
        ],
        speaker="Krótko przejść po czterech blokach; szczegóły na slajdach z ekranami.",
    )

    add_text_only_slide(
        prs,
        "Cel modułu mobilnego",
        [
            "Udostępnienie pracownikowi magazynu funkcji krytycznych poza stanowiskiem PC: przyjęcia, realizacja zadań.",
            "Ograniczenie błędów operacyjnych: identyfikacja kodem, filtrowanie katalogu, weryfikacja lokalizacji przy zamknięciu zadania.",
            "Spójność z warstwą serwerową: te same endpointy i reguły co aplikacja webowa — brak „drugiego systemu”.",
        ],
        speaker="Można powiązać z wymaganiami funkcjonalnymi pracy: mobilność, ergonomia, audyt.",
    )

    add_text_only_slide(
        prs,
        "Architektura integracji (skrót)",
        [
            "Klient Flutter — komunikacja HTTPS/JSON z API w wersji /api/v1.",
            "Uwierzytelnianie: token JWT (access); odświeżanie sesji zgodnie z polityką serwera.",
            "Encje kluczowe: dokument PZ (przyjęcie), pozycje z miejscem odłożenia, zadanie typu odłożenie (PUTAWAY), lokalizacje, stany.",
            "Zrzuty ekranu pochodzą z uruchomienia na macOS (ten sam kod co iOS/Android); dane z działającego backendu demonstracyjnego.",
        ],
        speaker="Jeśli promotor pyta o skalowanie: podkreślić bezstanowość API i możliwość wielu urządzeń.",
    )

    add_text_only_slide(
        prs,
        "Scenariusz demonstracyjny — przebieg",
        [
            "Operator loguje się na konto magazyniera (00002).",
            "Tworzy PZ: dostawca Northwind (SUP-001), towar PRD-1001, ilość, miejsce docelowe STO-01.",
            "Rejestruje przyjęcie — towar na buforze, generowane zadania odłożenia.",
            "Realizuje zadanie: start pracy → potwierdzenie kodem STO-01 → zadanie zakończone; weryfikacja na liście z ujawnieniem zakończonych.",
        ],
        speaker="To jest spójny, zamknięty scenariusz możliwy do powtórzenia na środowisku demo.",
    )

    for st in steps:
        add_screenshot_slide(
            prs,
            step=st["step"],
            title=st["title"],
            takeaway=st["takeaway"],
            bullets=st["bullets"],
            image=st["image"],
            speaker=st.get("speaker", ""),
        )

    add_text_only_slide(
        prs,
        "Podsumowanie",
        [
            "Zaprezentowano kompletny szlak: od utworzenia dokumentu PZ po zamknięcie zadania magazynowego z walidacją lokalizacji.",
            "Interfejs jest przygotowany pod pracę w warunkach magazynu (skan, zwięzłe etykiety, filtry).",
            "Moduł stanowi integralną część systemu WMS — wspólne API, role i model danych z aplikacją webową.",
            "Kierunki rozwoju (opcjonalnie do dyskusji): tryb offline z kolejką, więcej typów dokumentów mobilnych, telemetryka.",
        ],
        speaker="Zaproponować pytania promotora — przygotowane odpowiedzi: bezpieczeństwo JWT, przypisanie zadań, spójność stanów.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"OK: {OUT}")
    print(f"  Slajdów: {len(prs.slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
