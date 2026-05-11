#!/usr/bin/env python3
"""Generuje SCENARIUSZ_TESTOW_RECZNY_GIELDA_ZADAN_WMS.xlsx (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SCENARIUSZ_TESTOW_RECZNY_GIELDA_ZADAN_WMS.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WYNIK_OPTIONS = '"PASS,FAIL,CZĘŚCIOWO,N/D,BLOK"'


def main() -> None:
    wb = Workbook()

    # Instrukcja
    ws0 = wb.active
    ws0.title = "Instrukcja"
    lines = [
        ("Scenariusz testów ręcznych WMS — wersja po poprawce „giełda zadań”.", True),
        ("", False),
        ("Wynik wpisuj: PASS / FAIL / CZĘŚCIOWO / N/D / BLOK.", False),
        ("Do FAIL/BLOK dodaj krótki opis oraz dowód (screen/log).", False),
        ("", False),
        ("Klucz testu: zadania po PZ/MM/RW mają być NEW i widoczne dla WORKER.", True),
        ("Po „Rozpocznij pracę” zadanie ma się przypisać do WORKER.", True),
        ("", False),
        ("Dane organizacyjne:", True),
        ("Data testów:", False),
        ("Tester:", False),
        ("Środowisko (IP/URL):", False),
        ("Wersja / commit:", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = ws0.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 95

    ws = wb.create_sheet("Scenariusz", 1)
    headers = [
        "Obszar",
        "ID",
        "Krok",
        "Czynność",
        "Oczekiwany wynik",
        "Wynik",
        "Uwagi",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    rows = [
        ("M1 Logowanie", "M1.1", "Web ADMIN login", "Zaloguj 00001", "Logowanie PASS"),
        ("M1 Logowanie", "M1.2", "Mobilka WORKER login", "Zaloguj 00002", "Logowanie PASS"),
        ("M2 Giełda PZ", "M2.1", "Web ADMIN: utwórz PZ", "Nowy dokument PZ", "Dokument utworzony"),
        ("M2 Giełda PZ", "M2.2", "Web ADMIN: Zarejestruj", "PZ po rejestracji", "Status przechodzi, zadanie tworzy się"),
        ("M2 Giełda PZ", "M2.3", "Mobilka WORKER: lista zadań", "Otwórz zakładkę Zadania", "Widzisz PUTAWAY w statusie NEW"),
        ("M2 Giełda PZ", "M2.4", "Mobilka WORKER: start zadania", "Klik Rozpocznij pracę", "Zadanie przypisuje się do WORKER"),
        ("M2 Giełda PZ", "M2.5", "Mobilka WORKER: complete", "Zakończ zadanie", "Status COMPLETED"),
        ("M3 Giełda MM", "M3.1", "Web ADMIN: utwórz MM", "Nowy MM", "Dokument utworzony"),
        ("M3 Giełda MM", "M3.2", "Web ADMIN: zatwierdź MM", "Confirm", "Powstaje MOVE w NEW"),
        ("M3 Giełda MM", "M3.3", "Mobilka WORKER: start MOVE", "Rozpocznij", "MOVE przypisany do WORKER"),
        ("M4 Giełda RW", "M4.1", "Web ADMIN: utwórz RW", "Wybierz lokalizację i odbiorcę", "RW utworzony"),
        ("M4 Giełda RW", "M4.2", "Web ADMIN: zatwierdź/generuj", "Uruchom zadania", "PICKING w NEW"),
        ("M4 Giełda RW", "M4.3", "Mobilka WORKER: start PICKING", "Rozpocznij", "PICKING przypisany do WORKER"),
        ("M5 RBAC", "M5.1", "WORKER -> /users", "Spróbuj wejść", "Brak dostępu / redirect"),
        ("M5 RBAC", "M5.2", "WORKER cudze IN_PROGRESS", "Spróbuj zakończyć cudze", "HTTP 403 / blokada"),
        ("M6 Spójność", "M6.1", "Stock po zadaniach", "Sprawdź ilości", "Zmiany stanów zgodne"),
        ("M6 Spójność", "M6.2", "Ledger", "Sprawdź wpisy", "Wpisy dla operacji widoczne"),
        ("M6 Spójność", "M6.3", "Audit", "Sprawdź ślad", "Akcje użytkowników widoczne"),
        ("M7 Negatywny RW", "M7.1", "RW bez stanu", "Próba zatwierdzenia", "Czytelny błąd, brak ruchu"),
        ("M8 Responsywność", "M8.1", "Web 360px", "Podstawowe ekrany", "Brak krytycznych rozjazdów"),
        ("M8 Responsywność", "M8.2", "Web 1366px", "Desktop", "Układ poprawny"),
    ]

    row = 2
    for area, sid, step, action, expected in rows:
        vals = [area, sid, step, action, expected, None, None]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    widths = [20, 10, 28, 36, 36, 14, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{row-1}"

    dv = DataValidation(type="list", formula1=WYNIK_OPTIONS, allow_blank=True)
    dv.add(f"F2:F{row-1}")
    ws.add_data_validation(dv)

    ws2 = wb.create_sheet("Podsumowanie")
    summary = [
        ("Łączna liczba kroków", str(len(rows))),
        ("PASS", ""),
        ("FAIL", ""),
        ("CZĘŚCIOWO", ""),
        ("BLOK", ""),
        ("N/D", ""),
        ("", ""),
        ("Najważniejsze 3 wnioski", ""),
        ("1)", ""),
        ("2)", ""),
        ("3)", ""),
    ]
    for i, (k, v) in enumerate(summary, start=1):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v if v else None)
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 70

    wb.save(OUT)
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    main()
