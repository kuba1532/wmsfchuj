#!/usr/bin/env python3
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SPRAWOZDANIE_POSTEPU_DPZ_WMS_2026-05-11_PRO.docx"


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Strona ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run._r.append(fld)


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.runs[0].font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "D9E1F2")
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(10)
    doc.add_paragraph("")


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Strona tytułowa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DYPlOMOWY PROJEKT ZESPOŁOWY")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Sprawozdanie postępu prac")
    run.bold = True
    run.font.size = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("System WMS (Web + Mobilka Flutter)").font.size = Pt(16)

    for _ in range(6):
        doc.add_paragraph("")

    meta = [
        ("Projekt:", "System WMS"),
        ("Data:", date.today().strftime("%d.%m.%Y")),
        ("Zakres:", "Ryzyka, harmonogram, jakość, metodologia, narzędzia, stack"),
    ]
    for k, v in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{k} ")
        r1.bold = True
        p.add_run(v)

    doc.add_page_break()

    # Spis treści ręczny
    heading(doc, "Spis treści", 1)
    toc = [
        "1. Ryzyka projektowe wraz z planem mitygacji",
        "2. Harmonogram prac (wykres Gantta)",
        "3. Plan monitorowania jakości oraz formatka scenariusza testowego",
        "4. Metodologia pracy",
        "5. Wykorzystywane narzędzia",
        "6. Stack technologiczny rozwiązania",
        "7. Załączniki i artefakty",
        "8. Podsumowanie",
    ]
    for item in toc:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    heading(doc, "1. Ryzyka projektowe wraz z planem mitygacji", 1)
    doc.add_paragraph(
        "Poniższa tabela przedstawia kluczowe ryzyka projektu oraz działania ograniczające ich "
        "prawdopodobieństwo i wpływ. Rejestr ryzyk jest aktualizowany cyklicznie przed demo i release."
    )
    add_table(
        doc,
        ["ID", "Ryzyko", "Prawd.", "Wpływ", "Poziom", "Plan mitygacji", "Status"],
        [
            ["R1", "Niespójność danych web/mobilka", "Śr.", "Wys.", "Wys.", "Auto-refresh, odświeżanie po operacji, testy multi-device", "Monitorowane"],
            ["R2", "Błędy RW i walidacji lokalizacji", "Śr.", "Wys.", "Wys.", "Walidacje backend + testy negatywne + komunikaty", "Aktywne"],
            ["R3", "Luki RBAC w edge-case", "N/Śr.", "Wys.", "Śr.", "Regresja per rola + testy API 403", "Monitorowane"],
            ["R4", "Ryzyka iOS (provisioning/certyfikaty)", "Śr.", "Śr.", "Śr.", "Checklisty release i testy na fizycznym urządzeniu", "Aktywne"],
            ["R5", "Rozjazd oczekiwań interesariuszy", "Śr.", "Śr.", "Śr.", "Cykliczne synchronizacje, protokół decyzji", "Monitorowane"],
            ["R6", "Problemy środowiska demo (IP/CORS/sieć)", "Śr.", "Śr./Wys.", "Śr.", "Instrukcja DEMO + test dnia prezentacji", "Aktywne"],
        ],
    )

    heading(doc, "2. Harmonogram prac (wykres Gantta)", 1)
    doc.add_paragraph(
        "Harmonogram realizacji obejmuje iteracje marzec-maj 2026. Zastosowano podejście etapowe "
        "z domykaniem kamieni milowych i testów regresyjnych."
    )
    add_table(
        doc,
        ["Obszar", "Marzec", "Kwiecień", "Maj", "Status"],
        [
            ["Analiza wymagań i model procesu", "██", "–", "–", "Zakończone"],
            ["Architektura i model danych", "██", "▒▒", "–", "Zakończone"],
            ["Backend API", "██", "██", "▒▒", "Wysoka gotowość"],
            ["Frontend web", "██", "██", "▒▒", "Wysoka gotowość"],
            ["Mobilka Flutter", "–", "██", "██", "Finalizacja"],
            ["Testy i UAT", "–", "██", "██", "Aktywne"],
            ["Materiały oddaniowe", "–", "▒▒", "██", "Aktywne"],
        ],
    )
    add_table(
        doc,
        ["Milestone", "Plan", "Status"],
        [
            ["M1: MVP backend + auth", "Koniec marca", "Zamknięty"],
            ["M2: Procesy PZ/MM/RW web", "Połowa kwietnia", "Zamknięty"],
            ["M3: Integracja mobilki z API", "Koniec kwietnia", "Zamknięty"],
            ["M4: UAT i role", "Początek maja", "W toku"],
            ["M5: Pakiet oddaniowy", "11.05.2026", "W toku"],
        ],
    )

    heading(doc, "3. Plan monitorowania jakości oraz formatka scenariusza testowego", 1)
    doc.add_paragraph("Model kontroli jakości opiera się o cztery warstwy testów:")
    for bullet in [
        "Smoke techniczny (health, auth, krytyczne endpointy).",
        "Integracyjne testy API (dokumenty, zadania, reguły biznesowe).",
        "UAT per rola (ADMIN/MANAGER/FOREMAN/WORKER).",
        "Regresja UI web + mobilka z dowodami screenshot/log.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("")
    add_table(
        doc,
        ["KPI jakości", "Wartość docelowa"],
        [
            ["Krytyczne błędy 5xx", "0"],
            ["Krytyczne regresje flow PZ/MM/RW", "0"],
            ["Błędy autoryzacji poza oczekiwanym 403", "0"],
            ["Skuteczność wykonania scenariuszy UAT", ">=95%"],
        ],
    )
    doc.add_paragraph("Formatka scenariusza testowego (szablon):")
    add_table(
        doc,
        ["Pole", "Opis"],
        [
            ["ID scenariusza", "Unikalny identyfikator (np. UAT-RW-01)"],
            ["Obszar", "Auth / PZ / MM / RW / Inventory / RBAC"],
            ["Cel testu", "Zakres weryfikacji biznesowej i technicznej"],
            ["Precondition", "Dane wejściowe, konto, stan systemu"],
            ["Kroki wykonania", "Numerowana lista kroków testera"],
            ["Oczekiwany rezultat", "Efekt UI/API/DB"],
            ["Wynik", "PASS / FAIL / CZĘŚCIOWO / BLOK / N/D"],
            ["Dowód", "Zrzut, log, raport"],
            ["Uwagi", "Komentarz i rekomendacje"],
        ],
    )

    heading(doc, "4. Metodologia pracy", 1)
    doc.add_paragraph(
        "W projekcie zastosowano podejście Agile (Scrum/Kanban hybrid): krótkie iteracje, "
        "priorytetyzacja Must/Should/Could, regularne przeglądy i szybkie pętle feedbacku."
    )
    for bullet in [
        "Planowanie i rozbijanie backlogu na małe przyrosty.",
        "Code review i wspólna walidacja zmian.",
        "Testowanie krytycznych ścieżek po każdej większej zmianie.",
        "Zarządzanie ryzykiem i decyzjami na cotygodniowych syncach.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    heading(doc, "5. Wykorzystywane narzędzia", 1)
    add_table(
        doc,
        ["Obszar", "Narzędzia"],
        [
            ["Backend", "Python, FastAPI, SQLAlchemy, Alembic, Uvicorn"],
            ["Frontend", "React, TypeScript, Vite, MUI, React Query, Axios"],
            ["Mobilka", "Flutter, Dart, mobile_scanner, shared_preferences"],
            ["Baza danych", "MySQL, Docker Compose"],
            ["Testy", "pytest, vitest, testy E2E/UAT, dart analyze"],
            ["Jakość kodu", "Ruff, ESLint, TypeScript checks"],
            ["Współpraca", "Git, repo dokumentacyjne, przeglądy zmian"],
        ],
    )

    heading(doc, "6. Stack technologiczny rozwiązania", 1)
    doc.add_paragraph(
        "Architektura rozwiązania składa się z trzech warstw: klient web (SPA), klient mobilny "
        "oraz backend REST API, z centralną bazą danych MySQL i mechanizmem JWT + RBAC."
    )
    add_table(
        doc,
        ["Warstwa", "Technologie"],
        [
            ["Klient web", "React 19, TypeScript, Vite, MUI, Zod, React Hook Form"],
            ["Klient mobilny", "Flutter/Dart (iOS/Android), HTTP API, skaner kodów"],
            ["Backend API", "FastAPI, SQLAlchemy, Pydantic, JWT (access/refresh)"],
            ["Persistencja", "MySQL + migracje Alembic"],
            ["Autoryzacja", "RBAC per obszar funkcjonalny"],
        ],
    )

    heading(doc, "7. Załączniki i artefakty", 1)
    for item in [
        "RYZYKA_I_PLAN_MONITOROWANIA_WMS.xlsx",
        "SCENARIUS_TESTOW_RECZNY_WYMAGANIA.xlsx",
        "DEMO.md",
        "prezentacja_demo/ (materiały prezentacyjne)",
        "e2e_tests/ i wyniki_scenariusza_auto/ (raporty testowe)",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    heading(doc, "8. Podsumowanie", 1)
    doc.add_paragraph(
        "Projekt znajduje się na etapie finalizacji pakietu oddaniowego. Kluczowe procesy operacyjne "
        "(PZ/MM/RW), mechanizmy autoryzacji i część mobilna są zaimplementowane oraz testowane. "
        "Bieżący etap obejmuje domknięcie pełnej dokumentacji, walidację końcową i przygotowanie prezentacji."
    )

    # Stopka z numeracją
    footer_p = doc.sections[0].footer.paragraphs[0]
    set_page_number(footer_p)

    # Kolejne sekcje dziedziczą stopkę
    for section in doc.sections[1:]:
        section.footer.is_linked_to_previous = True

    doc.save(str(OUT))
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    build()
