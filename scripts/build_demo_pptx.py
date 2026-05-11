"""Generuje prezentację PPTX z 16 ekranów demo WMS.
Slajdy: tytuł + treść (kroki) po lewej, screenshot po prawej.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = Path("/Users/jakub/Downloads/WMS 2/prezentacja_demo")
SRC = ROOT / "zrzuty"
OUT = ROOT / "WMS_demo_prezentacja.pptx"

NAVY = RGBColor(0x0B, 0x1B, 0x2E)
BLUE = RGBColor(0x21, 0x96, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xCF, 0xD8, 0xDC)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xFF, 0x8F, 0x00)

SLIDES = [
    {
        "title": "WMS — Demo systemu",
        "subtitle": "Pełen przepływ: Przyjęcie → Rozmieszczenie → Przesunięcie → Wydanie",
        "image": None,
        "bullets": [
            "Backend: FastAPI + MySQL + SQLAlchemy + Alembic",
            "Frontend: React 19 + TypeScript + Material-UI + Vite",
            "Mobilka: Flutter (iOS + Android), skaner kodów",
            "Audit Log + Stock Ledger + JWT + zarządzanie rolami",
        ],
    },
    {
        "title": "1. Logowanie",
        "image": "01_login.png",
        "bullets": [
            "Krok: otwórz http://localhost:5173",
            "Kod logowania: 00001 (admin)",
            "Hasło: MojeHasloAdmina2026!WMS",
            "Klik: Zaloguj się",
        ],
        "speaker": "Logowanie 5-cyfrowym kodem operatora + hasło. JWT z auto-odświeżaniem sesji.",
    },
    {
        "title": "2. Pulpit (Dashboard)",
        "image": "02_dashboard.png",
        "bullets": [
            "Liczniki na żywo: produkty, zadania, dokumenty, blokady",
            "Auto-refresh co 2 minuty",
            "Z prawej: przełącznik motywu + wylogowanie",
        ],
        "speaker": "Dashboard pokazuje aktualny stan magazynu. „Zablokowany towar 3” to historyczne wpisy z fazy beta.",
    },
    {
        "title": "3. Lista PZ — Przyjęcia zewnętrzne",
        "image": "03_pz_lista.png",
        "bullets": [
            "Kliknij: Operacje → Przyjęcia (PZ) lub /documents/pz",
            "Widoczne wszystkie przyjęcia + status: Nowy / W trakcie / Zakończony",
            "Po prawej akcje wprost w wierszu: Zarejestruj / Rozmieść / Podgląd",
        ],
        "speaker": "Każdy dokument ma ścieżkę życia. Zielony = zakończony, pomarańczowy = w toku, szary = świeży.",
    },
    {
        "title": "4. Tworzenie nowego PZ",
        "image": "04_pz_form.png",
        "bullets": [
            "Klik: + Nowe przyjęcie",
            "Dostawca: SUP-001 — Northwind Supplies",
            "Produkt: PRD-0001 (Moduł podstawowy X)",
            "Odłożenie: STO-01 (już teraz wskazujesz, gdzie towar ma trafić)",
            "Ilość: 1 → Utwórz dokument",
        ],
        "speaker": "Kluczowa zmiana: w PZ od razu wskazujesz lokalizację docelową. To eliminuje „zgubiony towar na buforze”.",
    },
    {
        "title": "5. PZ utworzony — status „Nowy”",
        "image": "05_pz_lista_nowy.png",
        "bullets": [
            "Nowy dokument PZ/2026/025 trafia na górę listy",
            "Status: Nowy (jeszcze towar fizycznie nie przyjęty)",
            "W kolumnie Akcje: zielony przycisk „Zarejestruj”",
        ],
        "speaker": "Dokument istnieje, ale stock jeszcze nie został zaksięgowany. To stan oczekiwania na fizyczne potwierdzenie odbioru.",
    },
    {
        "title": "6. Klik „Zarejestruj” → towar na buforze",
        "image": "06_pz_w_trakcie.png",
        "bullets": [
            "Stock natychmiast zapisany na REC-01 jako AVAILABLE",
            "Wygenerowane zadanie PUTAWAY (REC-01 → STO-01)",
            "Status PZ: W trakcie (czeka na rozmieszczenie)",
        ],
        "speaker": "Jedno kliknięcie wykonuje zatwierdzenie + księgowanie + utworzenie zadania. To realistyczny scenariusz: dostawca przywiózł, ale jeszcze nie odłożono.",
    },
    {
        "title": "7. Lista zadań rozmieszczania",
        "image": "07_putaway_lista.png",
        "bullets": [
            "Kliknij „Rozmieść” lub przejdź do /putaway",
            "Pierwsze zadanie u góry: Moduł podstawowy X 1 szt REC-01 → STO-01",
            "Wcześniej zakończone zadania widoczne jako „Rozmieszczono”",
        ],
        "speaker": "Tu pracuje magazynier — widzi dokładnie co i skąd dokąd przenieść.",
    },
    {
        "title": "8. Stepper rozmieszczenia (3 kroki)",
        "image": "08_putaway_stepper.png",
        "bullets": [
            "Krok 1: Pobierz z bufora",
            "Krok 2: Przenieś na lokalizację",
            "Krok 3: Potwierdź rozmieszczenie",
            "Klik 3× „Następny krok / Zakończ”",
        ],
        "speaker": "Stepper wymusza dyscyplinę procesu — magazynier nie pominie kroku.",
    },
    {
        "title": "9. PZ zamknięty automatycznie",
        "image": "09_pz_zakonczony.png",
        "bullets": [
            "Po wykonaniu wszystkich zadań PUTAWAY status → Zakończony",
            "Stock przeniesiony z REC-01 na STO-01 (potwierdzone w Stanach)",
            "Cały ślad widoczny w Audit Log",
        ],
        "speaker": "Najważniejsze: dokument zamyka się sam, gdy fizyczna praca jest wykonana — nie przez „udajemy że zrobione”.",
    },
    {
        "title": "10. MM — Przesunięcia międzymagazynowe",
        "image": "10_mm_lista.png",
        "bullets": [
            "Kliknij: Operacje → Przesunięcia (MM)",
            "Tworzysz: z lokalizacji STO-A-01 → do PICK-A-01",
            "Krok: Utwórz → Zatwierdź → Wygeneruj zadania → Wykonaj",
        ],
        "speaker": "MM zasila strefę pickingową. Bez tego nie można pobrać towaru do wydania.",
    },
    {
        "title": "11. RW — Wydania",
        "image": "11_rw_lista.png",
        "bullets": [
            "Kliknij: Operacje → Wydania (RW)",
            "Odbiorca: Demo Klient, produkt + ilość",
            "Krok: Utwórz → Zatwierdź → Wygeneruj zadania → PICKING wykonany przez magazyniera",
        ],
        "speaker": "RW pobiera towar wyłącznie ze strefy pickingowej. To zabezpieczenie przed wydaniem czegoś co jest w głębi magazynu.",
    },
    {
        "title": "12. Stany magazynowe",
        "image": "12_stock.png",
        "bullets": [
            "Widok aktualnego stanu na każdej lokalizacji",
            "Statusy: Dostępny / Zablokowany / Zarezerwowany",
            "Ostrzeżenie u góry o pozycjach zablokowanych",
        ],
        "speaker": "Pełna transparentność: w każdej chwili widać ile czego i gdzie jest.",
    },
    {
        "title": "13. Rejestr ruchów (Stock Ledger)",
        "image": "13_ledger.png",
        "bullets": [
            "Każda zmiana stanu = wpis w ledgerze",
            "Przyjęcie / Rozmieszczenie / Wydanie / Przesunięcie",
            "Powiązanie z dokumentem (PZ/MM/RW)",
        ],
        "speaker": "To księga magazynowa. Niezmienialna historia ruchów — wymaganie audytowe.",
    },
    {
        "title": "14. Dziennik zdarzeń (Audit Log)",
        "image": "14_audit.png",
        "bullets": [
            "418 zarejestrowanych zdarzeń",
            "Kto, kiedy, jaką akcję, na jakim obiekcie",
            "Pełna ścieżka audytowa (utworzenie / zmiana / usunięcie / login)",
        ],
        "speaker": "Każdy klik użytkownika jest zapisywany. Można odtworzyć incydent.",
    },
    {
        "title": "15. Raporty (CSV / XLSX)",
        "image": "15_raporty.png",
        "bullets": [
            "4 raporty out-of-the-box: Stany / Historia / Realizacja / Analiza",
            "Liczniki rekordów z liveAPI",
            "Eksport do CSV i XLSX jednym klikiem",
        ],
        "speaker": "Eksporty potrzebne dla księgowości i kontroli. XLSX otwiera się w Excelu bez konwersji.",
    },
    {
        "title": "16. Użytkownicy + zaproszenia mailowe",
        "image": "16_users.png",
        "bullets": [
            "Lista 15 kont z rolami: Admin / Magazynier / Brygadzista / Kierownik",
            "+ Utwórz konto → mail z linkiem do ustawienia hasła",
            "5-cyfrowy kod logowania generowany automatycznie",
        ],
        "speaker": "Zarządzanie użytkownikami z RBAC. Nowy pracownik dostaje mail z linkiem — sam ustawia sobie hasło.",
    },
    {
        "title": "Podsumowanie",
        "image": None,
        "bullets": [
            "✓ PZ → bufor → PUTAWAY → automatyczne zamknięcie dokumentu",
            "✓ MM → potwierdzenie → MOVE task → przesunięcie międzylokalizacyjne",
            "✓ RW → potwierdzenie → PICKING task → wydanie towaru klientowi",
            "✓ Stock Ledger + Audit Log + raporty CSV/XLSX",
            "✓ Wieloplatformowość: web (desktop+mobile responsive) + Flutter (iOS+Android)",
            "✓ Bezpieczeństwo: JWT, role, hash haseł, link aktywacyjny mailem",
        ],
        "speaker": "Cały cykl od dostawcy do klienta jest zamknięty, audytowalny i wieloplatformowy.",
    },
]


def add_title_slide(prs, slide):
    layout = prs.slide_layouts[6]  # blank
    s = prs.slides.add_slide(layout)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY

    # Pasek akcent
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.4), prs.slide_width, Inches(0.05))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE

    title_box = s.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(12), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = slide["title"]
    r.font.size = Pt(54)
    r.font.bold = True
    r.font.color.rgb = WHITE

    sub_box = s.shapes.add_textbox(Inches(0.7), Inches(2.7), Inches(12), Inches(1.0))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = slide["subtitle"]
    r.font.size = Pt(22)
    r.font.color.rgb = GRAY

    # Bullets
    bul_box = s.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(12), Inches(3))
    tf = bul_box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(slide["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = "•  " + b
        r.font.size = Pt(20)
        r.font.color.rgb = WHITE
        p.space_after = Pt(8)


def add_content_slide(prs, slide):
    layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(layout)

    # Tło
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY

    # Tytuł
    title_box = s.shapes.add_textbox(Inches(0.4), Inches(0.25), prs.slide_width - Inches(0.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = slide["title"]
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = WHITE

    # Pasek akcent pod tytułem
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.05), Inches(0.8), Inches(0.06))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE

    # Bullets po lewej (połowa szerokości)
    bul_box = s.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(5.6), Inches(5.5))
    tf = bul_box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(slide["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + b
        r.font.size = Pt(16)
        r.font.color.rgb = WHITE
        p.space_after = Pt(6)

    # Notatki dla mówcy
    if slide.get("speaker"):
        s.notes_slide.notes_text_frame.text = slide["speaker"]

    # Screenshot po prawej
    if slide.get("image"):
        img_path = SRC / slide["image"]
        if img_path.exists():
            # Ramka pod obrazek
            frame = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(6.4), Inches(1.3),
                Inches(6.4), Inches(5.5),
            )
            frame.line.color.rgb = BLUE
            frame.line.width = Pt(1.5)
            frame.fill.solid()
            frame.fill.fore_color.rgb = RGBColor(0x05, 0x10, 0x1F)
            # Obrazek (zachowaj proporcje)
            s.shapes.add_picture(
                str(img_path),
                Inches(6.5), Inches(1.4),
                width=Inches(6.2),
            )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide in SLIDES:
        if slide.get("image") is None and "subtitle" in slide:
            add_title_slide(prs, slide)
        elif slide.get("image") is None:
            # Slajd końcowy bez obrazka — wyrenderuj z layoutem treści
            add_content_slide(prs, slide)
        else:
            add_content_slide(prs, slide)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"OK: {OUT}")
    print(f"  Slajdów: {len(prs.slides)}")


if __name__ == "__main__":
    main()
